"""Metadata-only NovPhy coverage and frozen-interface compatibility inventory."""
from __future__ import annotations

from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

from scripts.build_issue_45_evidence import CONSTRAINTS_WORKBOOK_REFERENCE
from tasks.task_generator.canonical_materialization import _author_scenario_object_ids, _read_template


SCENARIOS = {1: "single_force", 2: "multiple_forces", 3: "rolling", 4: "falling", 5: "sliding"}
NOVELTIES = {
    0: ("normal", "normal counterpart; may already contain a non-novel external force"),
    1: ("objects", "changed object appearance/class"),
    2: ("agents", "fan applies a horizontal force"),
    3: ("actions", "increased upward air-turbulence force"),
    4: ("interactions", "magnetic attraction/repulsion"),
    5: ("relations", "slingshot moves to the right"),
    6: ("environments", "inverted gravity"),
    7: ("goals", "air turbulence changes force direction"),
    8: ("events", "storm after first bird death"),
}
UPSTREAM = "https://github.com/phy-q/NovPhy/blob/main/README.md"


def workbook_records(path):
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = workbook["Task Variations"]
        result = {}
        for row_number, values in enumerate(sheet.iter_rows(values_only=True), 1):
            name = values[1]
            if not isinstance(name, str) or not name.startswith("00001_"):
                continue
            fields = name.split("_")
            family = "type"+fields[3]
            if family in result:
                raise ValueError(f"duplicate workbook family: {family}")
            result[family] = {"row": row_number, "novel_template_name": name,
                              "novelty_level": int(fields[4]), "scenario": int(fields[5]),
                              "normal_coordinates": list(values[2:8]), "novel_coordinates": list(values[8:14]),
                              "normal_restrictions": values[14], "novel_restrictions": values[15]}
        return result
    finally:
        workbook.close()


def inspect_template(path, template_root, vocabulary, constraints):
    relative = path.relative_to(template_root)
    level = int(relative.parts[0].removeprefix("novelty_level_")); family = relative.parts[1]
    scenario = int(family[-2:])
    if level not in NOVELTIES or scenario not in SCENARIOS:
        raise ValueError(f"unknown benchmark cell: {relative}")
    data = path.read_bytes()
    tree = _read_template(data)[0]
    _author_scenario_object_ids(tree)  # in-memory canonical authoring, no materialization or write
    objects = list(tree.find("GameObjects")); birds = list(tree.find("Birds"))
    slots = [n.attrib["scenarioObjectId"] for n in tree.iter() if "scenarioObjectId" in n.attrib]
    missing = sorted(set(slots)-set(vocabulary))
    counts = dict(sorted(Counter(s.split(":")[0] for s in slots).items()))
    sling = tree.find("Slingshot")
    sling_x = float(sling.attrib["x"])
    pigs = [float(n.attrib["x"]) for n in objects if n.tag == "Pig"]
    leftward = any(x < sling_x for x in pigs)
    extra = [{"tag":n.tag, "type":n.attrib.get("type"), "slot":n.attrib["scenarioObjectId"]}
             for n in objects if n.tag in ("Novelty", "ExternalAgent", "TNT")]
    binding = constraints.get(family)
    if binding is None or binding["scenario"] != scenario or (level and binding["novelty_level"] != level):
        raise ValueError(f"template/workbook counterpart mapping differs: {relative}")
    current = level == 0 and family in ("type010101", "type010102")
    issues = []
    if missing: issues.append("authored_slots_missing_from_frozen_CNN")
    if leftward: issues.append("leftward_launch_not_supported_by_current_negative_drag_x_grid")
    if not current: issues.append("outside_current_two_family_normal_collector_scope")
    if not current: issues.append("parser_generalization_and_generated_entity_inventory_not_validated")
    if level in (2, 3, 4, 6, 7, 8) or extra:
        issues.append("force_agent_event_state_observability_and_symbolic_semantics_not_validated")
    try:
        source_text = data.decode("utf-8")
    except UnicodeDecodeError:
        source_text = data.decode("utf-16")
    return {"template": str(relative), "family": family, "scenario": SCENARIOS[scenario],
            "scenario_index": scenario, "novelty_level": level, "novelty_category": NOVELTIES[level][0],
            "mechanism": NOVELTIES[level][1], "source_xml": source_text,
            "constraints": {"workbook": CONSTRAINTS_WORKBOOK_REFERENCE, "sheet":"Task Variations", **binding,
                            "active_coordinates": binding["novel_coordinates" if level else "normal_coordinates"],
                            "active_restrictions": binding["novel_restrictions" if level else "normal_restrictions"]},
            "authored_slots": slots, "slot_counts": counts, "missing_slots": missing,
            "object_types": dict(sorted(Counter(f"{n.tag}:{n.attrib.get('type','')}:{n.attrib.get('material','')}" for n in objects+birds).items())),
            "special_objects": extra, "camera": dict(tree.find("Camera").attrib),
            "slingshot": dict(sling.attrib), "current_source_scope": current,
            "static_slot_action_fit": not missing and not leftward,
            "runtime_validated_by_this_inventory": False, "issues": issues,
            "timing_requirements": "observe first bird death and later storm; a225-step shot may miss the event"
                if level == 8 else "actual launch/settlement and novelty effects require rendered development audit",
            "observability": "current runtime has image-derived slots and optional adjacent-frame motion, not true novelty IDs/forces/event memory",
            "prior_exposure": "family present in #62-#75 training/calibration/model_selection; lineages still distinct"
                if current else "outside #62-#75 family scope; global prior/final lineage exposure not audited"}


def build_inventory(repo_root, vocabulary):
    template_root = repo_root/"tasks/task_templates"
    constraints = workbook_records(repo_root/CONSTRAINTS_WORKBOOK_REFERENCE)
    paths = sorted(template_root.glob("novelty_level_*/type*/Levels/*.xml"))
    templates = [inspect_template(p, template_root, vocabulary, constraints) for p in paths]
    lookup = {(t["novelty_level"], t["family"]):t for t in templates}
    if len(templates) != 80 or len(lookup) != 80 or len(constraints) != 40:
        raise ValueError("expected the upstream80-template/40-counterpart inventory; review a versioned inventory change")
    pairs = []
    for family, row in sorted(constraints.items()):
        normal, novel = lookup[(0,family)], lookup[(row["novelty_level"],family)]
        pairs.append({"family":family, "scenario":normal["scenario"], "novelty_level":row["novelty_level"],
                      "normal":normal["template"], "novel":novel["template"], "workbook_row":row["row"],
                      "both_static_slot_action_fit":normal["static_slot_action_fit"] and novel["static_slot_action_fit"],
                      "pairing_basis":"upstream family and workbook linkage; not necessarily identical authored geometry",
                      "fresh_pair_cluster_rule":"same base scenario plus normal/novel variants stay in one partition and paired cluster"})
    cells = []
    for level in range(9):
        for scenario in range(1,6):
            members = [t for t in templates if t["novelty_level"] == level and t["scenario_index"] == scenario]
            expected = 8 if level == 0 else 1
            if len(members) != expected: raise ValueError("missing scenario/novelty cell")
            cells.append({"novelty_level":level,"novelty_category":NOVELTIES[level][0],"scenario":SCENARIOS[scenario],
                          "templates":[t["template"] for t in members],
                          "static_slot_action_fit_count":sum(t["static_slot_action_fit"] for t in members),
                          "current_source_count":sum(t["current_source_scope"] for t in members),
                          "runtime_or_performance_tested":False})
    # A template that does not fit does not make its whole scenario unsupported.
    recommended = []
    for scenario in ("rolling","falling","sliding"):
        candidates = [t for t in templates if t["novelty_level"] == 0 and t["scenario"] == scenario and t["static_slot_action_fit"]]
        if candidates: recommended.append(min(candidates,key=lambda t:t["template"]))
    recommended.extend(t for t in templates if (t["novelty_level"],t["family"]) in ((0,"type010102"),(1,"type010102")))
    return {"schema":"issue_76_novelty_inventory_v1", "upstream_reference":UPSTREAM,
            "metadata_only":True, "templates":templates, "pairs":pairs, "cells":cells,
            "vocabulary":vocabulary, "normal_templates":40, "novel_templates":40,
            "summary":{"cells":45,"pairs":40,"static_slot_action_fit_templates":sum(t["static_slot_action_fit"] for t in templates),
                       "missing_slot_templates":sum(bool(t["missing_slots"]) for t in templates),
                       "runtime_broader_coverage_validated":False},
            "wider_recommendation":{
                "status":"costed_proposal_not_authorized_or_executed",
                "candidate_templates":[{"template":t["template"],"static_slot_action_fit":t["static_slot_action_fit"],"issues":t["issues"]} for t in recommended],
                "selection":"lexicographically first static-compatible normal template per rolling/falling/sliding, plus matched type010102 appearance pair; no model outcomes",
                "first_engineering_step":"versioned condition-aware collector packaging and slot/observation audit; preserve existing collector",
                "planning_estimates_not_measurements":{"engineering_hours_cap":8,"development_smoke_captures_max":8,
                    "smoke_wall_minutes_cap":30,"smoke_fixed_steps_per_capture_max":600,"cpu_gib":3,"capture_output_gib":2},
                "requires_before_execution":["explicit approval for new compatibility work", "exact templates/generation seeds/roles freeze",
                    "rendered agent-frame/video audit with all attempts retained", "global lineage exposure audit", "no target-novelty fitting in zero-shot protocol"],
                "physical_novelty_next":"audit magnetic and inverted-gravity counterparts; missing slots/force-state requirements must be resolved symmetrically first",
                "adaptation_scope":"any CNN/carrier/action/history change and matched refit needs a separately costed approved plan; no refit budget silently granted"},
            "limitations":["template counts are not generated-runtime entity guarantees", "static fit is not parser recognition or legal trajectory reachability",
                           "no held-out outcomes or success-filter inventories inspected", "no global freshness claim", "no causal claim that normal-only tasks caused hybrid failure"],
            "fresh_evaluation_opened":False,"final_evaluation_opened":False,"issue_64_authorized":False}
